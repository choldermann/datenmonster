"""Inventur: Bestandsliste einfrieren, bewerten, Summen führen.

Ablauf einer Inventur:
  1. anlegen()   – Kopf mit Stichtag und Mandant
  2. befuellen() – führt das Bestands-Mapping aus und schreibt das Ergebnis als
                   Momentaufnahme in die Positionen. Ab hier ist die Liste
                   unabhängig von der WaWi.
  3. vorschlag_anwenden() – setzt Abwertungsvorschläge nach Restlaufzeit
  4. bewerten()  – der Anwender überschreibt einzelne Zeilen
  5. abschliessen() – friert den Lauf ein; er ist danach der Beleg.

Die Bestandsermittlung selbst steht bewusst NICHT hier, sondern in einem
Mapping (SQL bleibt SQL). Dieser Service kennt nur die Spaltennamen, die das
Mapping liefern muss – siehe SPALTEN.
"""
from datetime import date, datetime, timezone
from typing import Optional, List

from app.models.inventur import (InventurLauf, InventurPosition, STANDARD_STUFEN)

# Standardname des Mappings, das die Bestandsliste liefert. Kommt aus dem
# Template jtl_lager_cockpit; am Lauf überschreibbar.
STANDARD_MAPPING = "Inventur – Bestand zum Stichtag"

# Spalten, die das Bestands-Mapping liefern muss. Fehlt eine, bleibt das Feld
# leer – die Inventur läuft trotzdem, nur mit weniger Information.
SPALTEN = {
    "k_artikel": ("kArtikel", "k_artikel"),
    "c_artnr": ("ArtNr", "cArtNr", "Artikelnummer"),
    "artikelname": ("Artikel", "Artikelname", "cName"),
    "warengruppe": ("Warengruppe",),
    "hersteller": ("Hersteller",),
    "bestand": ("Bestand", "Menge"),
    "ek": ("EK", "EK netto"),
    "wert": ("Wert", "Bestandswert"),
    "mhd_frueh": ("MHD", "MHD früh", "MHD frühestes"),
    "resttage": ("Resttage", "Restzeit Tage"),
    "menge_abgelaufen": ("Menge abgelaufen",),
    "wert_abgelaufen": ("Wert abgelaufen",),
    "abgang_12m": ("Abgang 12M", "Abgang 12 Monate"),
    "reichweite_tage": ("Reichweite Tage", "Reichweite"),
    "letzter_abgang": ("Letzter Abgang",),
    "menge_ohne_partie": ("Menge ohne Charge", "Menge ohne Partie"),
    "chargen": ("Chargen",),
}


# Gründe für eine Mengenabweichung. Feste Auswahl statt nur Freitext: sonst lässt
# sich später nicht auswerten, wie viel Bruch, Schwund oder Fehlbuchung war.
DIFFERENZ_GRUENDE = {
    "bruch": "Bruch / Verderb",
    "schwund": "Schwund / Diebstahl",
    "fehlbuchung": "Fehlbuchung",
    "we_ungebucht": "Wareneingang nicht gebucht",
    "wa_ungebucht": "Warenausgang nicht gebucht",
    "fund": "Fund / Umlagerung",
    "sonstiges": "Sonstiges",          # braucht eine Notiz
}

# Unterhalb davon gilt eine Mengendifferenz als keine (Rundung aus der Abfrage).
MENGE_EPS = 0.0005


def _jetzt():
    return datetime.now(timezone.utc)


def _hole(row: dict, feld: str):
    """Liest einen Wert aus der Mapping-Zeile über die erlaubten Spaltennamen."""
    for name in SPALTEN.get(feld, ()):
        if name in row:
            return row[name]
    return None


def _zahl(v, default=0.0) -> float:
    if v is None or v == "":
        return default
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def _ganzzahl(v) -> Optional[int]:
    if v is None or v == "":
        return None
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None


def _datum(v) -> Optional[date]:
    """Nimmt date, datetime oder String. Deutsche Schreibweise inbegriffen –
    die Mappings formatieren MHD teilweise als TT.MM.JJJJ für die Anzeige."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


# ─── Bestandsliste aus dem Mapping holen ──────────────────────────────────────

def _resolve_mapping(db, project_id, name: str):
    from app.models.mapping import Mapping
    q = db.query(Mapping).filter(Mapping.name == name)
    if project_id is not None:
        q = q.filter(Mapping.project_id == project_id)
    return q.first()


# Eine Inventur kann viele Positionen haben – der Deckel muss großzügig sein,
# sonst fehlt am Jahresende stillschweigend Ware in der Liste.
ZEILEN_MAX = 20000


def bestandsliste_laden(db, lauf: InventurLauf, datum: Optional[date] = None) -> list:
    """Führt das Bestands-Mapping in der WaWi des Laufs aus.

    SQL bleibt SQL: Wie ein Stichtagsbestand ermittelt wird, steht in einem
    normalen Mapping und ist im Editor nachlesbar und anpassbar – dieselbe
    Trennung wie bei der Preisautomatik.
    """
    from app.services.mapping_service import MappingContext, execute_mapping
    from app.services import mandant_service

    name = lauf.quelle_mapping or STANDARD_MAPPING
    m = _resolve_mapping(db, lauf.project_id, name)
    if not m:
        raise ValueError(f"Die Auswertung „{name}“ ist in diesem Projekt nicht "
                         f"installiert.")
    ctx = MappingContext.from_orm(m)
    # Über lauf_vorbereiten, nicht mit einem selbstgebauten Parameter-Dict: die
    # Abfrage erwartet auch die Ausschlusslisten und die Schwellwerte des
    # Mandanten. Ein fehlender Listen-Parameter bricht den Lauf ab, ein falscher
    # Mandant wäre schlimmer – dann stünden fremde Bestände in der Inventur.
    ctx.run_params, _ = mandant_service.lauf_vorbereiten(
        # `datum` weicht nur für den Zähltag ab: dieselbe verprobte Rückrechnung,
        # nur auf einen anderen Tag.
        {"stichtag": datum or lauf.stichtag}, lauf.project_id, db,
        mandant_id=lauf.connection_id)
    mandant_service.verbindung_ersetzen(ctx, lauf.connection_id, db, m.project_id)
    if not ctx.targets:
        raise ValueError(f"Die Auswertung „{name}“ hat kein Ziel.")
    felder = ctx.targets[0].get("fields") or []
    res = execute_mapping(row_cap=ZEILEN_MAX, **ctx.to_execute_kwargs(felder, ZEILEN_MAX))
    fehler = res.get("errors") or []
    rows = res.get("rows") or []
    if fehler and not rows:
        raise ValueError(f"Bestandsermittlung fehlgeschlagen: {str(fehler[0])[:200]}")
    if rows and _hole(rows[0], "k_artikel") is None:
        raise ValueError("Die Bestandsauswertung liefert keine Spalte „kArtikel“ – "
                         "ohne sie lässt sich keine Position zuordnen.")
    return rows


# ─── Anlegen und Befüllen ─────────────────────────────────────────────────────

def anlegen(db, project_id: Optional[int], connection_id: int, name: str,
            stichtag: date, notiz: str = None, quelle_mapping: str = None,
            benutzer: str = None) -> InventurLauf:
    # Die Staffel der letzten Inventur desselben Mandanten übernehmen: wer einmal
    # „ab 60 Tagen drüber 100 %“ festgelegt hat, will das nicht jedes Jahr neu
    # eintragen. Kopiert, nicht verknüpft – die alte Inventur bleibt ihr Beleg.
    vorige_stufen = None
    for frueher in (db.query(InventurLauf)
                    .filter(InventurLauf.connection_id == connection_id)
                    .order_by(InventurLauf.stichtag.desc(), InventurLauf.id.desc())
                    .limit(20).all()):
        if frueher.abwertung_stufen:
            vorige_stufen = [dict(s) for s in frueher.abwertung_stufen]
            break
    lauf = InventurLauf(
        project_id=project_id,
        connection_id=connection_id,
        name=name,
        stichtag=stichtag,
        notiz=notiz,
        quelle_mapping=quelle_mapping or STANDARD_MAPPING,
        status="offen",
        erstellt_von=benutzer,
        abwertung_stufen=vorige_stufen,
    )
    _protokoll(lauf, "angelegt", benutzer, stichtag=stichtag.isoformat())
    db.add(lauf)
    db.commit()
    db.refresh(lauf)
    return lauf


def befuellen(db, lauf: InventurLauf, zeilen: List[dict], benutzer: str = None) -> dict:
    """Schreibt das Ergebnis des Bestands-Mappings als Momentaufnahme.

    Ein erneutes Befüllen ersetzt die Positionen vollständig – inklusive bereits
    erfasster Bewertungen. Das ist Absicht: eine halb aktualisierte Liste, in der
    Abwertungen an Mengen hängen, die es nicht mehr gibt, wäre schlimmer als
    ein bewusster Neuanfang. Der Aufrufer warnt vorher.
    """
    if lauf.status == "abgeschlossen":
        raise ValueError("Die Inventur ist abgeschlossen und kann nicht neu befüllt werden.")

    # Was dabei verloren geht, gehört in den Verlauf – wer später fragt, warum eine
    # Abwertung fehlt, soll es dort lesen können.
    verworfen = (db.query(InventurPosition)
                 .filter(InventurPosition.lauf_id == lauf.id,
                         InventurPosition.bewertung_art.isnot(None)).count())

    # Zählungen dagegen bleiben: sie kosten Stunden im Lager und hängen nicht an
    # der Bewertung. Wieder zugeordnet über Artikel und Charge.
    alte_zaehlungen = {}
    for p in (db.query(InventurPosition)
              .filter(InventurPosition.lauf_id == lauf.id,
                      InventurPosition.zaehlung_ebene.isnot(None)).all()):
        alte_zaehlungen[p.k_artikel] = {
            "ebene": p.zaehlung_ebene, "ist": p.ist_gezaehlt, "ohne_partie": p.ist_ohne_partie,
            "grund": p.differenz_grund, "notiz": p.differenz_notiz,
            "am": p.gezaehlt_am, "von": p.gezaehlt_von,
            "chargen": {_charge_schluessel(c): c.get("ist") for c in (p.chargen or [])
                        if c.get("ist") is not None},
        }

    db.query(InventurPosition).filter(InventurPosition.lauf_id == lauf.id).delete()

    ohne_ek = 0
    menge_ohne_partie_gesamt = 0.0
    angelegt = 0
    neue = []

    for row in zeilen or []:
        k_artikel = _ganzzahl(_hole(row, "k_artikel"))
        if k_artikel is None:
            continue
        bestand = _zahl(_hole(row, "bestand"))
        if abs(bestand) < 0.0001:
            continue
        ek = _zahl(_hole(row, "ek"))
        wert = _hole(row, "wert")
        wert = _zahl(wert) if wert is not None else bestand * ek
        if ek <= 0:
            ohne_ek += 1
        ohne_partie = _zahl(_hole(row, "menge_ohne_partie"))
        menge_ohne_partie_gesamt += ohne_partie

        chargen = _hole(row, "chargen")
        if isinstance(chargen, str):
            # Das Mapping kann die Partien als JSON-Text liefern (MSSQL FOR JSON).
            import json
            try:
                chargen = json.loads(chargen)
            except ValueError:
                chargen = []

        pos = InventurPosition(
            lauf_id=lauf.id,
            k_artikel=k_artikel,
            c_artnr=_hole(row, "c_artnr"),
            artikelname=_hole(row, "artikelname"),
            warengruppe=_hole(row, "warengruppe"),
            hersteller=_hole(row, "hersteller"),
            bestand=bestand,
            ek=ek,
            wert=wert,
            mhd_frueh=_datum(_hole(row, "mhd_frueh")),
            resttage=_ganzzahl(_hole(row, "resttage")),
            menge_abgelaufen=_zahl(_hole(row, "menge_abgelaufen")),
            wert_abgelaufen=_zahl(_hole(row, "wert_abgelaufen")),
            abgang_12m=_zahl(_hole(row, "abgang_12m")),
            reichweite_tage=_ganzzahl(_hole(row, "reichweite_tage")),
            letzter_abgang=_datum(_hole(row, "letzter_abgang")),
            menge_ohne_partie=ohne_partie,
            chargen=chargen if isinstance(chargen, list) else [],
        )
        db.add(pos)
        neue.append(pos)
        angelegt += 1

    hinweise = []
    if ohne_ek:
        hinweise.append({
            "art": "ohne_ek",
            # Diese Positionen stehen mit 0 € in der Liste – weder eine
            # Eingangsbuchung noch der Artikelstamm kennen einen Preis. Eine
            # Inventur muss das ausweisen, sonst fehlt der Wert unbemerkt.
            "text": f"{ohne_ek} Artikel haben überhaupt keinen Einkaufspreis "
                    f"(weder gebucht noch im Artikelstamm) und stehen deshalb "
                    f"mit 0 € in der Liste.",
            "anzahl": ohne_ek,
        })
    if menge_ohne_partie_gesamt > 0.0001:
        hinweise.append({
            "art": "ohne_partie",
            "text": f"{menge_ohne_partie_gesamt:,.0f} Stück lassen sich keiner "
                    f"Eingangspartie zuordnen (Altbestand aus der Zeit vor der "
                    f"Buchungshistorie) – für diese Menge gibt es kein MHD.".replace(",", "."),
            "anzahl": menge_ohne_partie_gesamt,
        })

    lauf.hinweise = hinweise

    # Mit Zähltag: Soll am Zähltag neu holen, bevor die alten Zählungen zurückkommen.
    db.flush()
    if lauf.zaehltag:
        _zaehltag_soll_eintragen(neue, bestandsliste_laden(db, lauf, datum=lauf.zaehltag),
                                 lauf, benutzer)
    uebernommen = 0
    for pos in neue:
        z = alte_zaehlungen.get(pos.k_artikel)
        if not z:
            continue
        _soll_sichern(pos)
        if z["ebene"] == "charge":
            pos.chargen = [{**c, "ist": z["chargen"].get(_charge_schluessel(c), c.get("ist"))}
                           for c in (pos.chargen or [])]
            pos.ist_ohne_partie = z["ohne_partie"]
        else:
            pos.ist_gezaehlt = z["ist"]
        pos.differenz_grund, pos.differenz_notiz = z["grund"], z["notiz"]
        pos.gezaehlt_am, pos.gezaehlt_von = z["am"], z["von"]
        _zaehlung_anwenden(pos, lauf.stichtag)
        uebernommen += 1

    _protokoll(lauf, "eingelesen", benutzer, positionen=angelegt,
               bewertungen_verworfen=verworfen or None,
               zaehlungen_uebernommen=uebernommen or None)
    lauf.updated_at = _jetzt()
    db.commit()
    summen_neu_rechnen(db, lauf)
    return {"positionen": angelegt, "hinweise": hinweise}


# ─── Bewerten ─────────────────────────────────────────────────────────────────

def _abwertung_rechnen(pos: InventurPosition, art: str, wert: float) -> tuple:
    """Rechnet aus der Eingabe des Anwenders (Prozent, neuer Stückwert oder
    Betrag) den Abwertungsbetrag und den neuen Positionswert.

    Ein Prozentsatz braucht einen BEZUG. „prozent" wirkt auf die ganze Position,
    „prozent_abgelaufen" nur auf den Wert der Partien, deren MHD am Stichtag
    vorbei war. Der Unterschied ist keine Kleinigkeit: Artikel 80123 bei PPS hat
    45.630 € Bestand, davon 31.936 € abgelaufen – die restlichen 13.694 € liegen
    in drei Chargen, die noch bis 2028 haltbar sind. „100 % Abschlag" auf die
    ganze Position hätte die mit abgeschrieben. Dieselbe Regel führt schon der
    Vorschlag (_vorschlag_fuer), nur staffelt der zusätzlich nach Restlaufzeit.

    Gedeckelt auf den Bestandswert: eine Abwertung kann eine Position auf null
    bringen, aber nicht ins Negative – ein negativer Lagerwert wäre keine
    Bewertung mehr, sondern ein Tippfehler mit Folgen für die Bilanz.
    """
    basis = pos.wert or 0.0
    if art == "prozent":
        betrag = basis * (wert / 100.0)
    elif art == "prozent_abgelaufen":
        betrag = (pos.wert_abgelaufen or 0.0) * (wert / 100.0)
    elif art == "stueckwert":
        betrag = basis - (pos.bestand or 0.0) * wert
    elif art == "betrag":
        betrag = wert
    else:
        raise ValueError(f"Unbekannte Bewertungsart: {art}")

    betrag = max(0.0, min(betrag, basis))
    return round(betrag, 2), round(basis - betrag, 2)


def bewerten(db, pos: InventurPosition, art: str, wert: float,
             grund: str = None, benutzer: str = None,
             vorschlag: bool = False) -> InventurPosition:
    betrag, neu = _abwertung_rechnen(pos, art, _zahl(wert))
    if not grund and art == "prozent_abgelaufen":
        # Der Bezug muss im Beleg stehen. Ohne ihn steht im Export des
        # Steuerberaters „50 %" neben einer Abwertung, die nur ein Drittel der
        # Position trifft – das sieht nach einem Rechenfehler aus.
        grund = (f"{_zahl(wert):g} % auf {_menge(pos.menge_abgelaufen)} Stück "
                 f"abgelaufene Ware ({_euro(pos.wert_abgelaufen)})")
    pos.bewertung_art = art
    pos.bewertung_wert = _zahl(wert)
    pos.abwertung_betrag = betrag
    pos.wert_neu = neu
    pos.grund = grund
    pos.vorschlag = vorschlag
    # Herkunft bleibt auch nach „Vorschläge übernehmen" erhalten – nur so kann eine
    # geänderte Staffel übernommene Staffel-Bewertungen neu rechnen, ohne echte
    # Handarbeit anzufassen.
    pos.bewertung_quelle = "staffel" if vorschlag else "hand"
    pos.bewertet_am = _jetzt()
    pos.bewertet_von = benutzer
    db.commit()
    return pos


def bewertung_loeschen(db, pos: InventurPosition) -> InventurPosition:
    pos.bewertung_art = None
    pos.bewertung_wert = None
    pos.abwertung_betrag = 0.0
    pos.wert_neu = None
    pos.grund = None
    pos.vorschlag = False
    pos.bewertung_quelle = None
    pos.bewertet_am = None
    pos.bewertet_von = None
    db.commit()
    return pos


def _euro(v) -> str:
    """Betrag in deutscher Schreibweise – landet in Begründungen und im Export."""
    return f"{_zahl(v):,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")


def _menge(v) -> str:
    """Mengen im Begründungstext ohne Nachkommastellen, wenn sie ganzzahlig sind."""
    n = _zahl(v)
    return f"{n:,.0f}".replace(",", ".") if abs(n - round(n)) < 0.001 else f"{n:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _vorschlag_fuer(pos: InventurPosition, stufen: List[dict],
                    stichtag: date) -> tuple:
    """Rechnet den Abwertungsvorschlag einer Position aus ihren Partien.

    Entscheidend: gestaffelt wird je CHARGE, nicht je Artikel. Ein Artikel mit
    zehn Chargen hat oft eine abgelaufene und neun frische — würde die Stufe
    „MHD überschritten → 100 %" den ganzen Positionswert treffen, wäre bei PPS
    eine Position mit 45.630 € Bestand komplett abgewertet worden, obwohl nur
    31.936 € davon tatsächlich abgelaufen sind.

    Gibt (betrag, grundtext) zurück; betrag 0 heißt „kein Vorschlag".
    """
    betrag = 0.0
    teile = {}   # label → Menge

    # Bei einer Zählung je ARTIKEL weiß niemand, welche Charge fehlt – die Partien
    # werden deshalb anteilig zur gezählten Menge gerechnet. Je Charge gezählt
    # tragen die Partien ihre gezählte Menge schon selbst.
    faktor = 1.0
    if pos.zaehlung_ebene == "artikel" and (pos.bestand_soll or 0.0) > MENGE_EPS:
        faktor = (pos.bestand or 0.0) / pos.bestand_soll

    for c in (pos.chargen or []):
        mhd = _datum(c.get("mhd"))
        if mhd is None:
            continue
        roh_menge = _zahl(c.get("menge"))
        # `wert` kommt exakt summiert aus der Abfrage; menge × ek wäre um die
        # Rundung des angezeigten EK daneben.
        wert = (_zahl(c.get("wert")) or roh_menge * _zahl(c.get("ek"))) * faktor
        menge = roh_menge * faktor
        if wert <= 0:
            continue
        resttage = (mhd - stichtag).days
        treffer = next((s for s in stufen if resttage <= _zahl(s.get("bis_tage"), 0)), None)
        if not treffer:
            continue
        prozent = _zahl(treffer.get("prozent"))
        if prozent <= 0:
            continue
        betrag += wert * prozent / 100.0
        label = treffer.get("label") or f"bis {_zahl(treffer.get('bis_tage')):g} Tage"
        teile[f"{label} ({prozent:g} %)"] = teile.get(f"{label} ({prozent:g} %)", 0) + menge

    if not teile:
        # Position ohne Chargenangabe, aber mit MHD am Artikel: dann trifft die
        # Stufe die ganze Position – mehr Information gibt es nicht her.
        if pos.resttage is None:
            return 0.0, None
        treffer = next((s for s in stufen if pos.resttage <= _zahl(s.get("bis_tage"), 0)), None)
        if not treffer:
            return 0.0, None
        prozent = _zahl(treffer.get("prozent"))
        if prozent <= 0:
            return 0.0, None
        return ((pos.wert or 0.0) * prozent / 100.0,
                f"{treffer.get('label') or 'Restlaufzeit'} ({prozent:g} %)")

    grund = "; ".join(f"{_menge(m)} Stück {label}" for label, m in teile.items())
    return betrag, grund


def _dauer(t: int) -> str:
    if t >= 30 and t % 30 == 0:
        m = t // 30
        return f"{m} {'Monat' if m == 1 else 'Monate'}"
    return f"{t} {'Tag' if t == 1 else 'Tage'}"


def _stufen_label(bis: int, vorher: Optional[int]) -> str:
    """Beschriftung einer Stufe aus ihren Resttagen und denen der Stufe davor.

    Dieselbe Regel wie `stufenLabel` in InventurStufenModal.tsx. Sie steht im
    Grundtext jeder vorgeschlagenen Bewertung und damit in der Liste für den
    Steuerberater – deshalb aus den Zahlen erzeugt statt frei eintippbar.
    Die Standardstaffel ergibt damit genau ihre bisherigen Texte.
    """
    if bis < 0:
        if vorher is None:
            return f"MHD mindestens {-bis} Tage überschritten"
        return f"MHD {-bis} bis {-vorher - 1} Tage überschritten"
    if bis == 0:
        if vorher is None:
            return "MHD überschritten"
        return f"MHD bis {-vorher - 1} Tage überschritten"
    if vorher is None:
        return f"Restlaufzeit bis {_dauer(bis)} (auch überschritten)"
    if vorher < 0:
        return f"Restlaufzeit bis {_dauer(bis)}"
    if vorher == 0:
        return f"unter {_dauer(bis)} Restlaufzeit"
    if vorher >= 30 and vorher % 30 == 0 and bis % 30 == 0:
        return f"{vorher // 30} bis {bis // 30} Monate Restlaufzeit"
    return f"{_dauer(vorher)} bis {_dauer(bis)} Restlaufzeit"


def stufen_pruefen(stufen) -> List[dict]:
    """Prüft und vereinheitlicht eine Staffel: ganze Resttage, 0–100 %, keine
    doppelten Resttage, aufsteigend sortiert, Beschriftung aus den Zahlen."""
    if not isinstance(stufen, list) or not stufen:
        raise ValueError("Die Staffel braucht mindestens eine Stufe.")
    if len(stufen) > 12:
        raise ValueError("Höchstens 12 Stufen.")
    sauber = []
    for s in stufen:
        try:
            bis_roh = float(s.get("bis_tage"))
            prozent = float(s.get("prozent"))
        except (TypeError, ValueError, AttributeError):
            raise ValueError("Jede Stufe braucht Resttage und einen Prozentsatz.")
        if bis_roh != int(bis_roh):
            raise ValueError("Resttage bitte als ganze Zahl.")
        bis = int(bis_roh)
        if not -3650 <= bis <= 3650:
            raise ValueError("Resttage müssen zwischen -3650 und 3650 liegen.")
        if not 0 <= prozent <= 100:
            raise ValueError("Der Prozentsatz muss zwischen 0 und 100 liegen.")
        sauber.append({"bis_tage": bis, "prozent": round(prozent, 2)})
    sauber.sort(key=lambda s: s["bis_tage"])
    tage = [s["bis_tage"] for s in sauber]
    if len(set(tage)) != len(tage):
        raise ValueError("Zwei Stufen haben dieselben Resttage.")
    vorher = None
    for s in sauber:
        s["label"] = _stufen_label(s["bis_tage"], vorher)
        vorher = s["bis_tage"]
    return sauber


def stufen_des_laufs(lauf: InventurLauf) -> List[dict]:
    """Die Staffel, nach der diese Inventur vorschlägt – eigene oder Standard."""
    return sorted(lauf.abwertung_stufen or STANDARD_STUFEN,
                  key=lambda s: _zahl(s.get("bis_tage"), 0))


def stufen_setzen(db, lauf: InventurLauf, stufen) -> InventurLauf:
    if lauf.status == "abgeschlossen":
        raise ValueError("Die Inventur ist abgeschlossen – ihre Staffel gehört zum Beleg.")
    lauf.abwertung_stufen = stufen_pruefen(stufen)
    lauf.updated_at = _jetzt()
    db.commit()
    db.refresh(lauf)
    return lauf


def vorschlag_anwenden(db, lauf: InventurLauf, stufen: List[dict] = None,
                       nur_unbewertete: bool = True, benutzer: str = None) -> dict:
    """Setzt Abwertungsvorschläge nach Restlaufzeit bis zum MHD.

    `nur_unbewertete` schützt die Handarbeit: wer eine Zeile schon selbst
    bewertet hat, will sie nicht von der Automatik überschrieben bekommen.
    Vorschläge sind als solche markiert (`vorschlag = True`) und lassen sich
    dadurch später von bestätigten Bewertungen unterscheiden.
    """
    if lauf.status == "abgeschlossen":
        raise ValueError("Die Inventur ist abgeschlossen.")

    stufen = stufen or stufen_des_laufs(lauf)
    # Aufsteigend: die engste Stufe (kleinste bis_tage) gewinnt.
    stufen = sorted(stufen, key=lambda s: _zahl(s.get("bis_tage"), 0))

    q = db.query(InventurPosition).filter(InventurPosition.lauf_id == lauf.id)
    gesetzt = entfernt = unveraendert = 0
    for pos in q.all():
        # Aus der Staffel stammt, was noch Vorschlag ist ODER übernommen wurde.
        # Nur echte Handbewertungen sind geschützt (NULL = Altbestand ohne
        # Herkunft, im Zweifel Handarbeit).
        aus_staffel = bool(pos.vorschlag) or pos.bewertung_quelle == "staffel"
        if nur_unbewertete and pos.bewertung_art and not aus_staffel:
            continue
        betrag, grund = _vorschlag_fuer(pos, stufen, lauf.stichtag)
        if betrag <= 0:
            # Eine Staffel-Bewertung, die die jetzige Staffel nicht mehr trägt, muss
            # weg – sonst bliebe die Abwertung der alten Staffel stehen und die Summe
            # stimmte mit keiner der beiden überein.
            if pos.bewertung_art and aus_staffel:
                bewertung_loeschen(db, pos)
                entfernt += 1
            continue
        if (pos.bewertung_art and aus_staffel and not pos.vorschlag
                and round(betrag, 2) == round(pos.abwertung_betrag or 0, 2)
                and grund == pos.grund):
            # Übernommen und unverändert: bleibt bestätigt. Nur was sich wirklich
            # ändert, wird wieder zum Vorschlag und muss neu übernommen werden.
            unveraendert += 1
            continue
        bewerten(db, pos, "betrag", round(betrag, 2), grund=grund,
                 benutzer=benutzer, vorschlag=True)
        gesetzt += 1

    summen_neu_rechnen(db, lauf)
    return {"vorschlaege": gesetzt, "entfernt": entfernt, "unveraendert": unveraendert}


def vorschlaege_bestaetigen(db, lauf: InventurLauf, benutzer: str = None) -> dict:
    """Macht aus Vorschlägen verbindliche Bewertungen (Häkchen unter die Liste)."""
    if lauf.status == "abgeschlossen":
        raise ValueError("Die Inventur ist abgeschlossen.")
    rows = (db.query(InventurPosition)
            .filter(InventurPosition.lauf_id == lauf.id,
                    InventurPosition.vorschlag.is_(True)).all())
    for pos in rows:
        pos.vorschlag = False
        pos.bewertet_von = benutzer or pos.bewertet_von
        pos.bewertet_am = _jetzt()
    db.commit()
    return {"bestaetigt": len(rows)}


# ─── Zählung (Soll/Ist) ───────────────────────────────────────────────────────

def _charge_schluessel(c: dict) -> tuple:
    """Charge + MHD + gerundeter EK – derselbe Schlüssel, nach dem die Abfrage die
    Partien bündelt. Damit finden sich Chargen aus zwei Läufen (Stichtag/Zähltag)
    oder vor und nach einem neuen Einlesen wieder."""
    return ((c.get("charge") or "").strip(), str(c.get("mhd") or "")[:10],
            round(_zahl(c.get("ek")), 4))


def _soll_sichern(pos: InventurPosition) -> None:
    """Hält die Buchmenge fest, bevor eine Zählung die gültigen Felder ändert.
    Einmalig: was schon gesichert ist, bleibt – sonst würde eine Zählung zum
    neuen Soll."""
    if pos.bestand_soll is None:
        pos.bestand_soll = pos.bestand or 0.0
        pos.wert_soll = pos.wert or 0.0
        pos.menge_abgelaufen_soll = pos.menge_abgelaufen or 0.0
        pos.wert_abgelaufen_soll = pos.wert_abgelaufen or 0.0
    if pos.menge_ohne_partie_soll is None:
        pos.menge_ohne_partie_soll = pos.menge_ohne_partie or 0.0
    chargen = []
    for c in (pos.chargen or []):
        c = dict(c)
        if "menge_soll" not in c:
            c["menge_soll"] = _zahl(c.get("menge"))
            c["wert_soll"] = (_zahl(c.get("wert")) if c.get("wert") is not None
                              else c["menge_soll"] * _zahl(c.get("ek")))
        chargen.append(c)
    pos.chargen = chargen


def _zaehlung_anwenden(pos: InventurPosition, stichtag: date) -> None:
    """Leitet aus Buchmenge und Zählung die gültigen Mengen und Werte ab (committet nicht).

    Differenz wird am ZÄHLTAG gemessen (gezählt − Soll am Zähltag) – dort steht der
    Zähler. Die gültige Menge zum Stichtag ist Soll zum Stichtag + Differenz: die
    Buchungen zwischen beiden Tagen heben sich damit heraus. Nie unter null.
    """
    _soll_sichern(pos)
    soll = pos.bestand_soll or 0.0
    wert_soll = pos.wert_soll or 0.0
    chargen = [dict(c) for c in (pos.chargen or [])]
    soll_zt = soll if pos.soll_zaehltag is None else pos.soll_zaehltag

    if any(c.get("ist") is not None for c in chargen) or pos.ist_ohne_partie is not None:
        pos.zaehlung_ebene = "charge"
        summe = 0.0
        for c in chargen:
            m_soll = _zahl(c.get("menge_soll"))
            c_soll_zt = m_soll if c.get("soll_zaehltag") is None else _zahl(c.get("soll_zaehltag"))
            if c.get("ist") is None:
                c["menge"], c["wert"] = m_soll, c.get("wert_soll")
                c.pop("differenz", None)
                continue
            d = _zahl(c["ist"]) - c_soll_zt
            c["differenz"] = round(d, 3)
            summe += d
            m = max(0.0, m_soll + d)
            c["menge"] = round(m, 3)
            c["wert"] = round(m * _zahl(c.get("ek")), 2)
        # Altbestand ohne Charge: eigene Zählzeile, sonst gilt seine Buchmenge.
        rest_wert_soll = wert_soll - sum(_zahl(c.get("wert_soll")) for c in chargen)
        rest_soll = pos.menge_ohne_partie_soll or 0.0
        rest_menge, rest_wert = rest_soll, rest_wert_soll
        if pos.ist_ohne_partie is not None:
            rest_soll_zt = rest_soll if pos.rest_soll_zaehltag is None else pos.rest_soll_zaehltag
            d = pos.ist_ohne_partie - rest_soll_zt
            summe += d
            rest_menge = max(0.0, rest_soll + d)
            rest_wert = (rest_wert_soll * rest_menge / rest_soll if rest_soll > MENGE_EPS
                         else rest_menge * (pos.ek or 0.0))
        pos.menge_ohne_partie = round(rest_menge, 3)
        pos.bestand = round(sum(_zahl(c["menge"]) for c in chargen) + rest_menge, 3)
        pos.wert = round(sum(_zahl(c["wert"]) for c in chargen) + rest_wert, 2)
        alt = [c for c in chargen if _datum(c.get("mhd")) and _datum(c.get("mhd")) < stichtag]
        pos.menge_abgelaufen = round(sum(_zahl(c["menge"]) for c in alt), 3)
        pos.wert_abgelaufen = round(sum(_zahl(c["wert"]) for c in alt), 2)
        pos.differenz = round(summe, 3)
        pos.ist_gezaehlt = round(soll_zt + summe, 3)
    elif pos.ist_gezaehlt is not None:
        pos.zaehlung_ebene = "artikel"
        d = pos.ist_gezaehlt - soll_zt
        m = max(0.0, soll + d)
        anteil = (m / soll) if soll > MENGE_EPS else None
        pos.differenz = round(d, 3)
        pos.bestand = round(m, 3)
        pos.wert = (round(wert_soll * anteil, 2) if anteil is not None
                    else round(m * (pos.ek or 0.0), 2))
        # Welche Charge fehlt, weiß man bei einer Artikelzählung nicht: die
        # abgelaufene Menge geht anteilig mit.
        pos.menge_abgelaufen = round((pos.menge_abgelaufen_soll or 0.0) * (anteil or 0.0), 3)
        pos.wert_abgelaufen = round((pos.wert_abgelaufen_soll or 0.0) * (anteil or 0.0), 2)
        for c in chargen:
            c["menge"], c["wert"] = c.get("menge_soll"), c.get("wert_soll")
            c.pop("differenz", None)
        pos.menge_ohne_partie = pos.menge_ohne_partie_soll or 0.0
    else:
        pos.zaehlung_ebene = None
        pos.differenz = None
        pos.bestand, pos.wert = soll, wert_soll
        pos.menge_abgelaufen = pos.menge_abgelaufen_soll or 0.0
        pos.wert_abgelaufen = pos.wert_abgelaufen_soll or 0.0
        for c in chargen:
            c["menge"], c["wert"] = c.get("menge_soll"), c.get("wert_soll")
            c.pop("differenz", None)
        pos.menge_ohne_partie = pos.menge_ohne_partie_soll or 0.0

    # Frühestes MHD der Ware, die (noch) da ist – eine auf 0 gezählte abgelaufene
    # Charge soll die Position nicht weiter rot färben.
    mhds = [_datum(c.get("mhd")) for c in chargen
            if _datum(c.get("mhd")) and _zahl(c.get("menge")) > MENGE_EPS]
    if mhds:
        pos.mhd_frueh = min(mhds)
        pos.resttage = (pos.mhd_frueh - stichtag).days
    pos.chargen = chargen


def _bewertung_nachziehen(pos: InventurPosition, lauf: InventurLauf, benutzer: str = None) -> None:
    """Nach geänderter Menge die Bewertung mitziehen (committet nicht).

    Handbewertung: dieselbe Eingabe (z.B. 50 %) auf die neue Menge. Staffel-
    Bewertung: neu vorschlagen; ein geänderter Betrag wird wieder zum Vorschlag."""
    if not pos.bewertung_art:
        return
    if pos.vorschlag or pos.bewertung_quelle == "staffel":
        betrag, grund = _vorschlag_fuer(pos, stufen_des_laufs(lauf), lauf.stichtag)
        if betrag <= 0:
            for feld in ("bewertung_art", "bewertung_wert", "wert_neu", "grund",
                         "bewertet_am", "bewertet_von", "bewertung_quelle"):
                setattr(pos, feld, None)
            pos.abwertung_betrag = 0.0
            pos.vorschlag = False
            return
        betrag = round(betrag, 2)
        if betrag != round(pos.abwertung_betrag or 0.0, 2) or grund != pos.grund:
            pos.bewertung_art, pos.bewertung_wert, pos.grund = "betrag", betrag, grund
            pos.vorschlag, pos.bewertung_quelle = True, "staffel"
            pos.bewertet_am, pos.bewertet_von = _jetzt(), benutzer
    betrag, neu = _abwertung_rechnen(pos, pos.bewertung_art, pos.bewertung_wert or 0.0)
    pos.abwertung_betrag, pos.wert_neu = betrag, neu


def zaehlung_setzen(db, lauf: InventurLauf, pos: InventurPosition,
                    felder: dict, benutzer: str = None) -> InventurPosition:
    """Speichert eine Zählung aus der Oberfläche (eine Position, sofort gespeichert)."""
    _zaehlung_felder(lauf, pos, felder, benutzer)
    db.commit()
    summen_neu_rechnen(db, lauf)
    return pos


def _zaehlung_felder(lauf: InventurLauf, pos: InventurPosition, felder: dict,
                     benutzer: str = None) -> None:
    """Wendet eine Zählung an, ohne zu speichern – der Import schreibt viele auf einmal.

    `felder` enthält nur, was sich ändert: `ist` (Artikel ohne Chargen, None =
    zurücksetzen), `charge_index` + `charge_ist` oder `chargen_ist` {index: menge},
    `ohne_partie_ist` (Altbestand ohne Charge), `grund`, `notiz`."""
    if lauf.status == "abgeschlossen":
        raise ValueError("Die Inventur ist abgeschlossen.")
    _soll_sichern(pos)
    mengen_geaendert = False

    if "charge_index" in felder:
        i = felder["charge_index"]
        chargen = [dict(c) for c in (pos.chargen or [])]
        if i is None or not 0 <= int(i) < len(chargen):
            raise ValueError("Charge nicht gefunden.")
        wert = felder.get("charge_ist")
        if wert is not None and wert < 0:
            raise ValueError("Eine gezählte Menge kann nicht negativ sein.")
        chargen[int(i)]["ist"] = None if wert is None else round(float(wert), 3)
        pos.chargen = chargen
        pos.ist_gezaehlt = None          # die Artikelmenge ergibt sich jetzt aus den Chargen
        mengen_geaendert = True

    if "chargen_ist" in felder:
        chargen = [dict(c) for c in (pos.chargen or [])]
        for i, wert in (felder["chargen_ist"] or {}).items():
            i = int(i)
            if not 0 <= i < len(chargen):
                raise ValueError("Charge nicht gefunden.")
            if wert is not None and wert < 0:
                raise ValueError("Eine gezählte Menge kann nicht negativ sein.")
            chargen[i]["ist"] = None if wert is None else round(float(wert), 3)
        pos.chargen = chargen
        pos.ist_gezaehlt = None
        mengen_geaendert = True

    if "ohne_partie_ist" in felder:
        wert = felder["ohne_partie_ist"]
        if wert is not None and wert < 0:
            raise ValueError("Eine gezählte Menge kann nicht negativ sein.")
        pos.ist_ohne_partie = None if wert is None else round(float(wert), 3)
        pos.ist_gezaehlt = None
        mengen_geaendert = True

    if "ist" in felder:
        wert = felder["ist"]
        if wert is not None and wert < 0:
            raise ValueError("Eine gezählte Menge kann nicht negativ sein.")
        if pos.chargen:
            # Wo Chargen existieren, wird je Charge gezählt – nur so bleibt
            # nachvollziehbar, welche Partie fehlt. Zurücksetzen geht trotzdem.
            if wert is not None:
                raise ValueError("Dieser Artikel hat Chargen und wird je Charge gezählt – "
                                 "bitte die Mengen an den Chargen eintragen.")
            pos.chargen = [{**c, "ist": None} for c in pos.chargen]
            pos.ist_ohne_partie = None
        pos.ist_gezaehlt = None if wert is None else round(float(wert), 3)
        mengen_geaendert = True

    if "grund" in felder:
        g = felder["grund"] or None
        if g and g not in DIFFERENZ_GRUENDE:
            raise ValueError("Unbekannter Grund.")
        pos.differenz_grund = g
    if "notiz" in felder:
        pos.differenz_notiz = (felder["notiz"] or "").strip() or None

    _zaehlung_anwenden(pos, lauf.stichtag)
    if pos.zaehlung_ebene is None:
        pos.differenz_grund = pos.differenz_notiz = None
        pos.gezaehlt_am = pos.gezaehlt_von = None
    else:
        if abs(pos.differenz or 0.0) <= MENGE_EPS:
            # Keine Abweichung, kein Grund – ein alter Grund wäre irreführend.
            pos.differenz_grund = pos.differenz_notiz = None
        if mengen_geaendert:
            pos.gezaehlt_am, pos.gezaehlt_von = _jetzt(), benutzer
    _bewertung_nachziehen(pos, lauf, benutzer)


def _zaehltag_soll_eintragen(positionen: List[InventurPosition], zeilen: List[dict],
                             lauf: InventurLauf, benutzer: str = None) -> None:
    """Trägt die Buchmengen am Zähltag an Positionen und Chargen ein (committet nicht).

    Chargen, die erst nach dem Stichtag eingelagert wurden, stehen am Zähltag im
    Regal – sie kommen mit Soll zum Stichtag 0 dazu, damit man sie mitzählen kann."""
    je_artikel, je_charge, je_rest = {}, {}, {}
    for row in zeilen or []:
        k = _ganzzahl(_hole(row, "k_artikel"))
        if k is None:
            continue
        je_artikel[k] = _zahl(_hole(row, "bestand"))
        je_rest[k] = _zahl(_hole(row, "menge_ohne_partie"))
        ch = _hole(row, "chargen")
        if isinstance(ch, str):
            import json
            try:
                ch = json.loads(ch)
            except ValueError:
                ch = []
        je_charge[k] = {_charge_schluessel(c): c for c in (ch if isinstance(ch, list) else [])}

    for pos in positionen:
        _soll_sichern(pos)
        chargen = [dict(c) for c in (pos.chargen or [])]
        if not zeilen:
            pos.soll_zaehltag = None
            pos.rest_soll_zaehltag = None
            chargen = [c for c in chargen if not c.get("nach_stichtag")]
            for c in chargen:
                c.pop("soll_zaehltag", None)
        else:
            pos.soll_zaehltag = round(je_artikel.get(pos.k_artikel, 0.0), 3)
            pos.rest_soll_zaehltag = round(je_rest.get(pos.k_artikel, 0.0), 3)
            zt = je_charge.get(pos.k_artikel, {})
            bekannt = set()
            for c in chargen:
                key = _charge_schluessel(c)
                bekannt.add(key)
                c["soll_zaehltag"] = round(_zahl(zt[key].get("menge")), 3) if key in zt else 0.0
            for key, c2 in zt.items():
                if key not in bekannt and _zahl(c2.get("menge")) > MENGE_EPS:
                    chargen.append({
                        "charge": c2.get("charge"), "mhd": c2.get("mhd"),
                        "ek": _zahl(c2.get("ek")), "einlagerungen": c2.get("einlagerungen"),
                        "menge": 0.0, "wert": 0.0, "menge_soll": 0.0, "wert_soll": 0.0,
                        "soll_zaehltag": round(_zahl(c2.get("menge")), 3),
                        "nach_stichtag": True,
                    })
        pos.chargen = chargen
        _zaehlung_anwenden(pos, lauf.stichtag)
        _bewertung_nachziehen(pos, lauf, benutzer)


def zaehltag_setzen(db, lauf: InventurLauf, zaehltag: Optional[date],
                    benutzer: str = None) -> dict:
    """Setzt den Zähltag und holt die Buchmengen dieses Tages (oder entfernt ihn)."""
    if lauf.status == "abgeschlossen":
        raise ValueError("Die Inventur ist abgeschlossen.")
    if zaehltag and zaehltag > date.today():
        raise ValueError("Der Zähltag kann nicht in der Zukunft liegen.")
    if zaehltag == lauf.stichtag:
        zaehltag = None
    zeilen = bestandsliste_laden(db, lauf, datum=zaehltag) if zaehltag else []
    positionen = db.query(InventurPosition).filter(InventurPosition.lauf_id == lauf.id).all()
    _zaehltag_soll_eintragen(positionen, zeilen, lauf, benutzer)
    lauf.zaehltag = zaehltag
    _protokoll(lauf, "zaehltag", benutzer,
               zaehltag=zaehltag.isoformat() if zaehltag else None)
    db.commit()
    summen_neu_rechnen(db, lauf)
    return {"positionen": len(positionen)}


# ─── Zählliste: ausdrucken und zurückspielen ──────────────────────────────────

ZAEHLLISTE_BLATT = "Zählliste"


def _schluessel_text(art: str, pos: InventurPosition, c: dict = None) -> str:
    """Stabiler Zeilenschlüssel der Zählliste. Artikel + Charge + MHD + EK statt
    Datenbank-ID: so passt eine Liste auch nach „Bestände neu einlesen" noch."""
    if art == "C":
        charge, mhd, ek = _charge_schluessel(c)
        return f"C|{pos.k_artikel}|{charge}|{mhd}|{ek:.4f}"
    return f"{art}|{pos.k_artikel}"


def zaehlliste_zeilen(db, lauf: InventurLauf, mit_soll: bool = False) -> List[dict]:
    """Zeilen der Zählliste: je Charge, je Altbestand ohne Charge, je Artikel ohne
    Chargen. Sortiert nach Warengruppe und Artikel, Chargen nach MHD – was im Regal
    zusammen liegt, steht auch auf der Liste zusammen. Soll ist die Buchmenge am
    Zähltag (sonst zum Stichtag), nur wenn gewünscht (sonst Blindzählung)."""
    mit_zaehltag = lauf.zaehltag is not None
    positionen = db.query(InventurPosition).filter(InventurPosition.lauf_id == lauf.id).all()
    positionen.sort(key=lambda p: ((p.warengruppe or "").lower(),
                                   (p.artikelname or "").lower(), p.c_artnr or ""))
    zeilen = []
    for p in positionen:
        basis = {"Warengruppe": p.warengruppe, "Artikelnummer": p.c_artnr,
                 "Artikel": p.artikelname}
        if p.chargen:
            for c in sorted(p.chargen, key=lambda c: (_datum(c.get("mhd")) or date.max,
                                                      c.get("charge") or "")):
                soll = c.get("menge_soll", c.get("menge"))
                if mit_zaehltag and c.get("soll_zaehltag") is not None:
                    soll = c.get("soll_zaehltag")
                # Am Zähltag schon verbraucht: nichts zu zählen.
                if _zahl(soll) <= MENGE_EPS and c.get("ist") is None:
                    continue
                z = {**basis, "Schlüssel": _schluessel_text("C", p, c),
                     "Charge": (c.get("charge") or "(ohne Nr.)") + ("  [neu]" if c.get("nach_stichtag") else ""),
                     "MHD": _datum(c.get("mhd"))}
                if mit_soll:
                    z["Soll"] = round(_zahl(soll), 3)
                zeilen.append(z)
            rest = (p.menge_ohne_partie_soll if p.menge_ohne_partie_soll is not None
                    else p.menge_ohne_partie)
            if mit_zaehltag and p.rest_soll_zaehltag is not None:
                rest = p.rest_soll_zaehltag
            if _zahl(rest) > MENGE_EPS or p.ist_ohne_partie is not None:
                z = {**basis, "Schlüssel": _schluessel_text("R", p),
                     "Charge": "ohne Charge (Altbestand)", "MHD": None}
                if mit_soll:
                    z["Soll"] = round(_zahl(rest), 3)
                zeilen.append(z)
        else:
            soll = p.bestand_soll if p.bestand_soll is not None else p.bestand
            if mit_zaehltag and p.soll_zaehltag is not None:
                soll = p.soll_zaehltag
            z = {**basis, "Schlüssel": _schluessel_text("A", p), "Charge": "",
                 "MHD": p.mhd_frueh}
            if mit_soll:
                z["Soll"] = round(_zahl(soll), 3)
            zeilen.append(z)
    return zeilen


def _import_zahl(v) -> Optional[float]:
    """Menge aus einer Excel-Zelle: Zahl, oder Text mit deutschem Komma („3,5")."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def zaehlliste_importieren(db, lauf: InventurLauf, inhalt: bytes, benutzer: str = None,
                           dateiname: str = None) -> dict:
    """Spielt eine ausgefüllte Zählliste zurück.

    Leere Ist-Zellen ändern nichts – so können mehrere Teams Teillisten
    zurückspielen. Grund und Notiz gehören zur Position; mehrere Notizen eines
    Artikels werden aneinandergehängt. Alles in einem Durchgang, eine Summenrechnung."""
    if lauf.status == "abgeschlossen":
        raise ValueError("Die Inventur ist abgeschlossen.")
    import io
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(inhalt), data_only=True, read_only=True)
    except Exception:
        raise ValueError("Die Datei ist keine lesbare Excel-Datei (.xlsx).")
    ws = wb[ZAEHLLISTE_BLATT] if ZAEHLLISTE_BLATT in wb.sheetnames else wb.worksheets[0]
    reihen = ws.iter_rows(values_only=True)
    kopf = [str(k or "").strip() for k in (next(reihen, None) or [])]

    def spalte(name):
        return kopf.index(name) if name in kopf else None
    i_key, i_ist, i_grund, i_notiz = spalte("Schlüssel"), spalte("Ist"), spalte("Grund"), spalte("Notiz")
    if i_key is None or i_ist is None:
        raise ValueError("Das ist keine Zählliste – die Spalten „Schlüssel“ und „Ist“ fehlen.")

    def zelle(row, i):
        return row[i] if i is not None and i < len(row) else None

    gruende = {v.lower(): k for k, v in DIFFERENZ_GRUENDE.items()}
    gruende.update({k: k for k in DIFFERENZ_GRUENDE})
    positionen = {p.k_artikel: p for p in
                  db.query(InventurPosition).filter(InventurPosition.lauf_id == lauf.id).all()}
    je_pos, fehler, fremd = {}, [], []
    zeilen = leer = werte = 0

    for nr, row in enumerate(reihen, start=2):
        schluessel = str(zelle(row, i_key) or "").strip()
        if not schluessel:
            continue
        zeilen += 1
        ist_roh = zelle(row, i_ist)
        if ist_roh is None or str(ist_roh).strip() == "":
            leer += 1
            continue
        ist = _import_zahl(ist_roh)
        if ist is None or ist < 0:
            fehler.append(f"Zeile {nr}: „{ist_roh}“ ist keine gültige Menge")
            continue
        try:
            art, rest = schluessel.split("|", 1)
            if art == "C":
                k_txt, rest = rest.split("|", 1)
                charge, mhd, ek = rest.rsplit("|", 2)
            else:
                k_txt = rest
            pos = positionen.get(int(k_txt))
        except (ValueError, TypeError):
            fremd.append(f"Zeile {nr}")
            continue
        if pos is None:
            fremd.append(f"Zeile {nr}")
            continue
        felder = je_pos.setdefault(pos.k_artikel, {})
        if art == "A":
            felder["ist"] = ist
        elif art == "R":
            felder["ohne_partie_ist"] = ist
        elif art == "C":
            ziel = (charge, mhd, round(_zahl(ek), 4))
            idx = next((i for i, c in enumerate(pos.chargen or [])
                        if _charge_schluessel(c) == ziel), None)
            if idx is None:
                fremd.append(f"Zeile {nr} ({pos.c_artnr}, Charge {charge})")
                continue
            felder.setdefault("chargen_ist", {})[idx] = ist
        else:
            fremd.append(f"Zeile {nr}")
            continue
        werte += 1

        g = zelle(row, i_grund)
        if g is not None and str(g).strip():
            gid = gruende.get(str(g).strip().lower())
            if gid:
                felder["grund"] = gid
            else:
                fehler.append(f"Zeile {nr}: Grund „{g}“ unbekannt")
        n = zelle(row, i_notiz)
        if n is not None and str(n).strip():
            felder["notiz"] = "; ".join(x for x in (felder.get("notiz"), str(n).strip()) if x)

    uebernommen = 0
    for k, felder in je_pos.items():
        pos = positionen[k]
        try:
            _zaehlung_felder(lauf, pos, felder, benutzer)
            uebernommen += 1
        except ValueError as e:
            fehler.append(f"{pos.c_artnr or k}: {e}")

    _protokoll(lauf, "zaehlliste", benutzer, datei=dateiname, werte=werte,
               positionen=uebernommen, fehler=len(fehler) or None)
    db.commit()
    summen_neu_rechnen(db, lauf)
    return {"zeilen": zeilen, "leer": leer, "werte": werte, "positionen": uebernommen,
            "nicht_zugeordnet": fremd[:10], "nicht_zugeordnet_anzahl": len(fremd),
            "fehler": fehler[:10], "fehler_anzahl": len(fehler)}


def _abweichungen_ohne_grund(db, lauf: InventurLauf) -> int:
    from sqlalchemy import func as sqlfunc, or_, and_
    return (db.query(InventurPosition)
            .filter(InventurPosition.lauf_id == lauf.id,
                    InventurPosition.zaehlung_ebene.isnot(None),
                    sqlfunc.abs(InventurPosition.differenz) > MENGE_EPS,
                    or_(InventurPosition.differenz_grund.is_(None),
                        and_(InventurPosition.differenz_grund == "sonstiges",
                             InventurPosition.differenz_notiz.is_(None))))
            .count())


# ─── Summen und Abschluss ─────────────────────────────────────────────────────

def summen_neu_rechnen(db, lauf: InventurLauf) -> InventurLauf:
    from sqlalchemy import func as sqlfunc
    q = db.query(
        sqlfunc.count(InventurPosition.id),
        sqlfunc.coalesce(sqlfunc.sum(InventurPosition.wert), 0.0),
        sqlfunc.coalesce(sqlfunc.sum(InventurPosition.abwertung_betrag), 0.0),
        sqlfunc.count(InventurPosition.bewertung_art),
    ).filter(InventurPosition.lauf_id == lauf.id).one()

    lauf.positionen_anzahl = int(q[0] or 0)
    lauf.bestand_wert = round(float(q[1] or 0.0), 2)
    lauf.abwertung_summe = round(float(q[2] or 0.0), 2)
    lauf.wert_nach_abwertung = round(lauf.bestand_wert - lauf.abwertung_summe, 2)
    lauf.bewertete_positionen = int(q[3] or 0)
    z = db.query(
        sqlfunc.count(InventurPosition.zaehlung_ebene),
        sqlfunc.coalesce(sqlfunc.sum(InventurPosition.wert
                                     - sqlfunc.coalesce(InventurPosition.wert_soll,
                                                        InventurPosition.wert)), 0.0),
    ).filter(InventurPosition.lauf_id == lauf.id).one()
    lauf.gezaehlte_positionen = int(z[0] or 0)
    lauf.differenz_wert = round(float(z[1] or 0.0), 2)
    lauf.abweichungen_ohne_grund = _abweichungen_ohne_grund(db, lauf)
    lauf.updated_at = _jetzt()
    db.commit()
    db.refresh(lauf)
    return lauf


def abschliessen(db, lauf: InventurLauf, benutzer: str = None) -> InventurLauf:
    """Friert den Lauf ein. Offene Vorschläge gelten damit als bestätigt –
    wer abschließt, steht für die Zahlen ein."""
    if lauf.status == "abgeschlossen":
        return lauf
    ohne_grund = _abweichungen_ohne_grund(db, lauf)
    if ohne_grund:
        raise ValueError(f"{ohne_grund} Mengenabweichung{'en haben' if ohne_grund != 1 else ' hat'} "
                         "noch keinen Grund (bei „Sonstiges“ eine Notiz). Ohne Grund lässt "
                         "sich die Inventur nicht abschließen.")
    res = vorschlaege_bestaetigen(db, lauf, benutzer)
    summen_neu_rechnen(db, lauf)
    lauf.status = "abgeschlossen"
    lauf.abgeschlossen_am = _jetzt()
    lauf.abgeschlossen_von = benutzer
    _protokoll(lauf, "abgeschlossen", benutzer,
               vorschlaege_bestaetigt=res.get("bestaetigt") or None,
               abwertung=lauf.abwertung_summe,
               wert_nach_abwertung=lauf.wert_nach_abwertung,
               gezaehlt=lauf.gezaehlte_positionen or None,
               differenz_wert=lauf.differenz_wert if lauf.gezaehlte_positionen else None)
    db.commit()
    db.refresh(lauf)
    return lauf


def _protokoll(lauf: InventurLauf, aktion: str, benutzer: str = None, **daten) -> None:
    """Hängt einen Eintrag an den Verlauf der Inventur (committet nicht selbst).

    Neue Liste statt append: eine JSON-Spalte bemerkt Änderungen an der
    vorhandenen Liste nicht und würde den Eintrag still verwerfen."""
    eintrag = {"aktion": aktion, "am": _jetzt().isoformat(timespec="seconds"), "von": benutzer}
    eintrag.update({k: v for k, v in daten.items() if v is not None})
    lauf.protokoll = list(lauf.protokoll or []) + [eintrag]


def wieder_oeffnen(db, lauf: InventurLauf, benutzer: str = None) -> InventurLauf:
    """Nur für den Irrtumsfall. Der Abschluss bleibt in `abgeschlossen_am`
    stehen, damit sichtbar bleibt, dass der Beleg einmal fertig war – und das
    Öffnen selbst steht im Verlauf."""
    if lauf.status != "abgeschlossen":
        return lauf
    _protokoll(lauf, "wieder_geoeffnet", benutzer,
               abschluss_vom=lauf.abgeschlossen_am.isoformat(timespec="seconds")
               if lauf.abgeschlossen_am else None)
    lauf.status = "offen"
    lauf.updated_at = _jetzt()
    db.commit()
    db.refresh(lauf)
    return lauf


# ─── Ausgabe für den Steuerberater ────────────────────────────────────────────

# Für den Steuerberater lesbar statt der internen Kennung – „prozent_abgelaufen"
# sagt ihm nichts, „% auf abgelaufene Chargen" schon.
ART_LABEL = {
    "prozent": "% auf Position",
    "prozent_abgelaufen": "% auf abgelaufene Chargen",
    "stueckwert": "neuer Stückwert",
    "betrag": "Abwertungsbetrag",
}

EXPORT_SPALTEN = [
    ("c_artnr", "Artikelnummer"),
    ("artikelname", "Artikel"),
    ("warengruppe", "Warengruppe"),
    ("hersteller", "Hersteller"),
    ("bestand_soll", "Menge Soll"),
    ("bestand", "Menge Ist"),
    ("differenz_stichtag", "Differenz"),
    ("gezaehlt", "gezählt"),
    ("ek", "EK netto"),
    ("wert", "Wert zum EK"),
    ("differenz_wert", "Differenzwert"),
    ("differenz_grund", "Differenzgrund"),
    ("differenz_notiz", "Notiz Differenz"),
    ("mhd_frueh", "MHD"),
    ("resttage", "Resttage"),
    ("menge_abgelaufen", "Menge abgelaufen"),
    ("abgang_12m", "Abgang 12 Monate"),
    ("reichweite_tage", "Reichweite Tage"),
    ("bewertung_art", "Bewertungsart"),
    ("bewertung_wert", "Bewertungswert"),
    ("abwertung_betrag", "Abwertung"),
    ("wert_neu", "Wert nach Abwertung"),
    ("grund", "Grund"),
]

# Excel-Formate je Spalte. Mengen bleiben „Standard": ein Format wie '#,##0' würde
# eine Bruchmenge (Meter, Liter) in der Anzeige stumm runden.
_EURO = '#,##0.00 "€"'
EXPORT_FORMATE = {
    "EK netto": '#,##0.00## "€"',
    "Wert zum EK": _EURO,
    "Differenzwert": _EURO,
    "MHD": "DD.MM.YYYY",
    "Resttage": "#,##0",
    "Reichweite Tage": "#,##0",
    "Bewertungswert": "#,##0.00##",
    "Abwertung": _EURO,
    "Wert nach Abwertung": _EURO,
}
EXPORT_SUMMEN = ["Wert zum EK", "Differenzwert", "Abwertung", "Wert nach Abwertung"]


def export_zeilen(db, lauf: InventurLauf, datum_als_text: bool = True) -> List[dict]:
    """Flache Zeilen für CSV/XLSX – bewusst ohne Chargen-Detail, das würde die
    Liste für den Steuerberater vervielfachen. Die Partien stehen im Drilldown.

    `datum_als_text=False` lässt Datumswerte als Datum stehen – für Excel, damit
    die MHD-Spalte richtig sortiert und nach Zeitraum filterbar ist."""
    rows = (db.query(InventurPosition)
            .filter(InventurPosition.lauf_id == lauf.id)
            .order_by(InventurPosition.wert.desc()).all())
    out = []
    for p in rows:
        zeile = {}
        for feld, label in EXPORT_SPALTEN:
            v = getattr(p, feld, None)
            if isinstance(v, date) and datum_als_text:
                v = v.strftime("%d.%m.%Y")
            if feld == "bewertung_art" and v:
                v = ART_LABEL.get(v, v)
            # Zählung: Soll/Differenz immer zum STICHTAG – das ist die Liste des
            # Steuerberaters. Ungezählt gilt die Buchmenge (Differenz leer).
            gezaehlt = p.zaehlung_ebene is not None
            if feld == "bestand_soll" and v is None:
                v = p.bestand
            if feld == "differenz_stichtag":
                v = (round((p.bestand or 0.0) - (p.bestand_soll if p.bestand_soll is not None
                                                 else (p.bestand or 0.0)), 3) if gezaehlt else None)
            if feld == "differenz_wert":
                v = (round((p.wert or 0.0) - (p.wert_soll if p.wert_soll is not None
                                              else (p.wert or 0.0)), 2) if gezaehlt else None)
            if feld == "gezaehlt":
                v = ("je Charge" if p.zaehlung_ebene == "charge" else "ja") if gezaehlt else "nein"
            if feld == "differenz_grund" and v:
                v = DIFFERENZ_GRUENDE.get(v, v)
            # Unbewertete Positionen behalten ihren vollen Wert – sonst summiert
            # der Steuerberater eine Spalte mit Löchern.
            if feld == "wert_neu" and v is None:
                v = p.wert
            zeile[label] = v
        out.append(zeile)
    return out
